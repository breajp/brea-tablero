"""Lee la planilla de Brea y deja datos/tablero.json: operaciones, participaciones y lo que hay que revisar.

La hoja maestra es "2026". Las hojas por agente son copias hechas a mano: no se leen como fuente,
solo se comparan contra la maestra para mostrar dónde no coinciden.
"""
import json
import os
import re
import sys
import unicodedata
from datetime import date, datetime
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.utils import column_index_from_string as col, get_column_letter

# El runner de GitHub está en UTC: de 21 a 24 h de Argentina, date.today() ya es mañana.
AHORA = datetime.now(ZoneInfo("America/Argentina/Buenos_Aires"))
HOY = AHORA.date()
# Problemas de forma de la planilla (encabezados repetidos, costos sin fecha, usuarios mal cargados):
# van a la pantalla Revisar del admin como kind "planilla".
AVISOS = []
ESTADOS = {
    "A": ("activa", "Activa"),
    "B": ("reservada", "Reservada"),
    "C": ("firma", "Firma pautada"),
    "D": ("vendida", "Vendida"),
    "E": ("alquilada", "Alquilada"),
    "F": ("comprada", "Comprada"),
    "G": ("pausada", "Pausada"),
    "H": ("baja", "Baja"),
}
DINERO = ["AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ"]
NO_SON_COPIAS = {"2026", "COMISIONES", "Hoja 6"}
# Las columnas se buscan por su encabezado (fila 2): en octubre de 2026 Agente y Productor
# cambiaron de lugar, y leer por letra cruzaba los roles sin dar ningún error.
ENCABEZADOS = {
    "estado": "Estado", "propiedad": "Propiedad", "productor": "Productor", "agente": "Agente",
    "lider": "Lider", "operacion": "Operación", "tipologia": "Tipología",
    "vendedor": "Acuerdo Vendedor", "comprador": "Acuerdo comprador", "interno": "Acuerdo Interno (Agente)",
    "data": "Data", "ingreso": "Fecha Ingreso", "reserva": "Fecha Reserva o baja", "conformacion": "Fecha Conformación",
    "impacto": "Mes Impacto", "publi": "$ Publi", "cierre": "$ Cierrre", "nego": "% NEGO", "escritura": "$ Escritura",
    "bruto": "Bruto General", "bruto_productor": "Bruto Productor", "bruto_agente": "Bruto Agente",
    "productor_neto": "$ Productor", "agente_neto": "$ Agente", "oficina": "OFICINA", "lider_neto": "$ Lider",
    "gerencia": "Gerencia", "brea": "Brea", "firma": "Fecha y Hora",
}


def sin_tildes(texto):
    return unicodedata.normalize("NFD", str(texto)).encode("ascii", "ignore").decode()


def slug(texto):
    return re.sub(r"[^a-z0-9]+", "-", sin_tildes(texto).lower()).strip("-")


class Directorio:
    """Quién es quién, desde la pestaña Personas: un nombre nuevo o cambiado llega al tablero sin tocar código."""

    def __init__(self, wb):
        self.alias, self.nombres, self.tasas, self.roles, self.desconocidos = {}, {}, {}, {}, {}
        ws = wb["Personas"]
        hdr = {str(ws.cell(1, i).value).strip(): i for i in range(1, ws.max_column + 1) if ws.cell(1, i).value}
        for r in range(2, ws.max_row + 1):
            planilla, nombre = ws.cell(r, hdr["Nombre en la planilla"]).value, ws.cell(r, hdr["Persona"]).value
            if not planilla or not nombre:
                continue
            pid = slug(nombre)
            self.alias[sin_tildes(planilla).strip().lower()] = pid
            self.alias.setdefault(sin_tildes(nombre).strip().lower(), pid)
            self.nombres.setdefault(pid, str(nombre).strip())
            tasa = ws.cell(r, hdr["% como productor"]).value
            if isinstance(tasa, (int, float)):
                self.tasas[pid] = tasa
            if "Rol" in hdr and ws.cell(r, hdr["Rol"]).value:
                self.roles[pid] = str(ws.cell(r, hdr["Rol"]).value).strip()

    def persona(self, valor, productor=None):
        if not valor or not str(valor).strip():
            return None
        clave = sin_tildes(valor).strip().lower()
        if clave in self.alias:
            return self.alias[clave]
        # "Gonza" como agente es el Gonza de la unidad del productor de esa fila (3A o CABA).
        if productor and productor.startswith(slug(valor)):
            return productor
        pid = slug(valor)
        self.nombres.setdefault(pid, str(valor).strip())
        self.desconocidos.setdefault(pid, str(valor).strip())
        return pid

    def conocida(self, valor):
        return bool(valor) and sin_tildes(valor).strip().lower() in self.alias

    def gerente(self):
        return next((p for p, rol in self.roles.items() if rol.lower() == "gerente"), None)


def encabezados(ws, fila, hoja):
    """La columna de cada encabezado. Si uno se repite vale el primero (el de más a la izquierda),
    sin importar en qué orden se recorran, y se avisa: en Costos hubo un "Mes" en A y otro en H, y el de H dejaba los costos vacíos."""
    hdr, repetidos = {}, []
    for i in range(1, ws.max_column + 1):
        v = ws.cell(fila, i).value
        if v is None or not str(v).strip():
            continue
        h = str(v).strip()
        if h in hdr:
            repetidos.append(f"{h} ({get_column_letter(hdr[h])} y {get_column_letter(i)})")
        else:
            hdr[h] = i
    if repetidos:
        AVISOS.append({"kind": "planilla", "row": 0, "property": f"Pestaña {hoja}",
                       "detail": f"Encabezados repetidos: {', '.join(repetidos)}. Se lee la primera columna de cada uno."})
    return hdr


def tabla(wb, hoja, fila_encabezado=1, necesarias=()):
    """Las filas de una pestaña como diccionarios por encabezado; las filas vacías se saltean."""
    if hoja not in wb.sheetnames:
        if necesarias:
            AVISOS.append({"kind": "planilla", "row": 0, "property": f"Pestaña {hoja}", "detail": "No está en la planilla: ¿se renombró?"})
        return []
    ws = wb[hoja]
    hdr = encabezados(ws, fila_encabezado, hoja)
    faltan = [h for h in necesarias if h not in hdr]
    if faltan:
        AVISOS.append({"kind": "planilla", "row": 0, "property": f"Pestaña {hoja}",
                       "detail": f"En la fila {fila_encabezado} no están las columnas {', '.join(faltan)}: ¿se renombró un encabezado o se agregó una fila arriba? Esa parte del tablero queda vacía."})
    filas = []
    for r in range(fila_encabezado + 1, ws.max_row + 1):
        fila = {h: ws.cell(r, i).value for h, i in hdr.items()}
        if any(v not in (None, "") for v in fila.values()):
            fila["_fila"] = r
            filas.append(fila)
    return filas


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def plata(x):
    return f"{x:,.0f}".replace(",", ".")


def formula_legible(patron, fila):
    # La planilla está en es_ES: una fórmula se lee con ";". Un valor escrito a mano (188,99) queda como está.
    texto = patron.replace("#", str(fila))
    return texto.replace(",", ";") if texto.startswith("=") else texto


def sin_parentesis_de_mas(formula):
    # "=((VLOOKUP(…))*AC#)" calcula lo mismo que "=(VLOOKUP(…))*AC#": un paréntesis que envuelve toda
    # la fórmula no la cambia, y avisarlo como "fórmula cambiada a mano" es un falso aviso.
    if not formula.startswith("="):
        return formula
    cuerpo = formula[1:].replace(" ", "")
    while cuerpo.startswith("(") and cuerpo.endswith(")"):
        nivel = 0
        for i, ch in enumerate(cuerpo):
            nivel += ch == "("
            nivel -= ch == ")"
            if nivel == 0:
                break
        if i != len(cuerpo) - 1:
            break
        cuerpo = cuerpo[1:-1]
    return "=" + cuerpo


def pct(v):
    # Los honorarios son fracciones chicas (0,015, 0,0415): redondear a centavos los rompe.
    return round(float(v), 6) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def dia(v):
    return v.date().isoformat() if isinstance(v, datetime) else None


def mes(v):
    # "Mes Impacto" se carga como 3/3, 4/4, 5/5: el día repite el mes para que la celda muestre "3- marzo".
    return f"{v.year}-{v.month:02d}" if isinstance(v, datetime) else None


def leer_categorias(ruta):
    """La tabla de categorías de la planilla de comisiones de Gonzalo (Hoja1, filas 4 a 7)."""
    ws = openpyxl.load_workbook(ruta, data_only=True).worksheets[0]
    nombres = {"JUNIOR": "Junior", "ACTIVO": "Activo", "SENIOR": "Senior", "EXPERTO": "Experto"}
    categorias = []
    for r in range(4, 8):
        texto = str(ws.cell(r, 3).value or "").upper()
        nombre = next((v for k, v in nombres.items() if k in texto), None)
        agente = ws.cell(r, 6).value
        if not nombre or not isinstance(agente, (int, float)):
            continue
        meta = ws.cell(r, 9).value
        categorias.append({
            "name": nombre, "agentRate": round(agente, 4), "leaderRate": round(0.65 - agente, 4),
            "breaRate": round(ws.cell(r, 8).value or 0.35, 4),
            "goal": meta if isinstance(meta, (int, float)) else None,
            "requirement": requisito_legible(ws.cell(r, 4).value),
        })
    return categorias


def categorias_de_la_planilla(wb):
    """La pestaña Categorías de la planilla ordenada, si existe: es la misma tabla, con el acuerdo de Flor."""
    if "Categorías" not in wb.sheetnames:
        return []
    ws = wb["Categorías"]
    categorias = []
    for r in range(2, ws.max_row + 1):
        nombre, agente = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not nombre or not isinstance(agente, (int, float)):
            break
        meta = ws.cell(r, 5).value
        categorias.append({"name": str(nombre).strip(), "agentRate": round(agente, 4), "leaderRate": round(ws.cell(r, 3).value or 0, 4),
                           "breaRate": round(ws.cell(r, 4).value or 0, 4), "goal": meta if isinstance(meta, (int, float)) and meta else None,
                           "requirement": requisito_legible(ws.cell(r, 6).value) if isinstance(meta, (int, float)) and meta else None})
    return categorias


def requisito_legible(texto):
    # "2 AÑOS DESDE EL INGRESO EN ESTA CATEGORÍA y LLEGAR A LA META $$ x año" → "2 años en esta categoría y llegar a la meta en el año"
    if not texto:
        return None
    t = re.sub(r"\s+", " ", str(texto)).strip().lower()
    t = t.replace("desde el ingreso en esta categoría", "en esta categoría").replace(" $$ x año", " en el año").replace(" de $$", "")
    return t


def categoria_de(tasa, categorias):
    if not isinstance(tasa, (int, float)):
        return None
    return next((c["name"] for c in categorias if abs(c["agentRate"] - tasa) < 0.001), None)


def control_contra_la_planilla(wb, operaciones, participaciones, costos, nombres):
    """Los totales que calcula la planilla (Producción y Resumen mensual) contra los del tablero.
    Si no coinciden, alguno de los dos lee distinto una fila: se avisa en vez de mostrar dos verdades."""
    avisos = []
    cerca = lambda a, b: abs((a or 0) - (b or 0)) <= 1
    if "Producción" in wb.sheetnames:
        ws = wb["Producción"]
        anio, hdr = ws["B1"].value, {str(ws.cell(3, i).value).strip(): i for i in range(1, ws.max_column + 1) if ws.cell(3, i).value}
        if isinstance(anio, (int, float)) and "Producción bruta (US$)" in hdr:
            for r in range(4, ws.max_row + 1):
                nombre = ws.cell(r, 1).value
                if not nombre or str(nombre).strip().lower() == "total":
                    continue
                pid = slug(nombre)
                hoja = ws.cell(r, hdr["Producción bruta (US$)"]).value
                tablero = sum(s["gross"] for s in participaciones
                              if s["person"] == pid and s["stage"] == "cerrada" and (s["month"] or "").startswith(str(int(anio))))
                if isinstance(hoja, (int, float)) and not cerca(hoja, tablero):
                    avisos.append({"kind": "planilla", "row": 0, "property": f"Producción · {nombres.get(pid, nombre)}",
                                   "detail": f"Producción bruta {int(anio)}: US$ {plata(hoja)} en la pestaña Producción y US$ {plata(tablero)} en el tablero."})
    if "Resumen mensual" in wb.sheetnames:
        ws = wb["Resumen mensual"]
        hdr = {str(ws.cell(3, i).value).strip(): i for i in range(1, ws.max_column + 1) if ws.cell(3, i).value}
        columnas = {"Facturación bruta (US$)": "gross", "Neto Brea (US$)": "brea"}
        for r in range(4, ws.max_row + 1):
            cuando = mes(ws.cell(r, 1).value)
            if not cuando:
                continue
            cerradas = [o for o in operaciones if o["status"] in ("vendida", "alquilada", "comprada")
                        and o["billingMonth"] == cuando and str(o["producerLabel"] or "").strip().lower() != "gonza 3a"]
            for titulo, campo in columnas.items():
                hoja = ws.cell(r, hdr[titulo]).value if titulo in hdr else None
                tablero = sum(o[campo] or 0 for o in cerradas)
                if isinstance(hoja, (int, float)) and not cerca(hoja, tablero):
                    avisos.append({"kind": "planilla", "row": 0, "property": f"Resumen mensual · {cuando}",
                                   "detail": f"{titulo}: US$ {plata(hoja)} en la planilla y US$ {plata(tablero)} en el tablero."})
            hoja = ws.cell(r, hdr["Costos (US$)"]).value if "Costos (US$)" in hdr else None
            if isinstance(hoja, (int, float)) and not cerca(hoja, costos.get(cuando, 0)):
                avisos.append({"kind": "planilla", "row": 0, "property": f"Resumen mensual · {cuando}",
                               "detail": f"Costos: US$ {plata(hoja)} en la planilla y US$ {plata(costos.get(cuando, 0))} en el tablero."})
    return avisos


def main(ruta, salida, ruta_categorias=None):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    wf = openpyxl.load_workbook(ruta)
    ws, fs = wb["2026"], wf["2026"]
    letras = {h: get_column_letter(i) for h, i in encabezados(ws, 2, "2026").items()}
    faltan = [h for h in ENCABEZADOS.values() if h not in letras]
    if faltan:
        raise SystemExit(f"La hoja 2026 no tiene estas columnas: {faltan}")
    L = {k: letras[h] for k, h in ENCABEZADOS.items()}
    # "Mes de cobro" (AR) es la columna que usan Participaciones, Producción y Resumen mensual:
    # leerla hace que el tablero y la planilla imputen cada operación al mismo mes.
    if "Mes de cobro" in letras:
        L["cobro"] = letras["Mes de cobro"]
    c = lambda r, k: ws.cell(r, col(L.get(k, k))).value

    directorio = Directorio(wb)
    tasas, NOMBRES = directorio.tasas, directorio.nombres
    gerente = directorio.gerente()

    operaciones, participaciones, revisar = [], [], []
    # Explicado por Gonzalo (22/10): Tres Arroyos es una unidad aparte, sin gerencia, y lo que
    # parece faltar ahí es a propósito. Las bolsas ya no se usan: sus fórmulas son historia.
    tres_arroyos, bolsas = set(), set()
    grupos = {}
    for r in range(3, ws.max_row + 1):
        prop = c(r, "propiedad")
        estado = str(c(r, "estado") or "")[:1]
        if not prop or estado not in ESTADOS:
            continue
        estado_id, _ = ESTADOS[estado]
        listado, reserva = dia(c(r, "ingreso")), dia(c(r, "reserva"))
        publi, cierre = num(c(r, "publi")), num(c(r, "cierre"))
        nego = num(c(r, "nego"))
        cobro = mes(c(r, "cobro")) if "cobro" in L else (mes(c(r, "impacto")) or mes(c(r, "conformacion")))
        if not cobro and estado not in "DEF":
            # Lo cerrado va solo por Mes de cobro, como en Producción. Lo que no se cobró (en curso,
            # pausada, baja) se sigue ubicando en el mes de la reserva o la baja, como hasta ahora.
            cobro = mes(c(r, "reserva"))
        nota = str(c(r, "data")).strip() if c(r, "data") else None
        if str(c(r, "productor") or "").strip().lower() == "gonza 3a":
            tres_arroyos.add(r)
        if nota and "bolsa" in nota.lower():
            bolsas.add(r)

        dias = None
        if listado:
            hasta = date.fromisoformat(reserva) if reserva else HOY
            dias = (hasta - date.fromisoformat(listado)).days

        # Una operación en bolsa se carga en dos filas, una por inmobiliaria o agente: misma
        # propiedad, mismo precio y misma fecha. Para contar operaciones y volumen va una sola vez.
        clave = (re.sub(r"\W", "", prop.lower()), cierre, reserva) if cierre else ("fila", r)
        trato = grupos.setdefault(clave, f"op-{len(grupos) + 1}")

        op = {
            "row": r, "deal": trato, "property": prop.strip(), "status": estado_id,
            "producer": directorio.persona(c(r, "productor")), "producerLabel": c(r, "productor"),
            "agent": directorio.persona(c(r, "agente"), directorio.persona(c(r, "productor"))), "leader": directorio.persona(c(r, "lider")),
            "kind": c(r, "operacion"), "typology": c(r, "tipologia"),
            "sellerFee": pct(c(r, "vendedor")), "buyerFee": pct(c(r, "comprador")), "internalSplit": pct(c(r, "interno")),
            "notes": nota, "listedAt": listado, "reservedAt": reserva, "signedAt": dia(c(r, "firma")),
            "billingMonth": cobro, "listPrice": publi, "closePrice": cierre, "deedPrice": num(c(r, "escritura")),
            "negotiationPct": nego, "daysListed": dias,
            "gross": num(c(r, "bruto")), "producerGross": num(c(r, "bruto_productor")), "agentGross": num(c(r, "bruto_agente")),
            "producerNet": num(c(r, "productor_neto")), "agentNet": num(c(r, "agente_neto")), "office": num(c(r, "oficina")),
            "leaderNet": num(c(r, "lider_neto")), "management": num(c(r, "gerencia")), "brea": num(c(r, "brea")),
            "potential": round(publi * c(r, "vendedor"), 2) if publi and isinstance(c(r, "vendedor"), (int, float)) and c(r, "vendedor") > 0 else None,
        }
        operaciones.append(op)

        if estado in "BCDEF":
            for rol, quien, bruto, neto in (
                ("productor", op["producer"], op["producerGross"], op["producerNet"]),
                ("agente", op["agent"], op["agentGross"], op["agentNet"]),
                ("lider", op["leader"], op["leaderNet"], op["leaderNet"]),
                # La gerencia no es producción (bruto 0), pero es plata que cobra quien tiene rol Gerente.
                ("gerencia", gerente, 0, op["management"]),
            ):
                if quien and (bruto or neto):
                    participaciones.append({
                        "row": r, "deal": trato, "person": quien, "role": rol,
                        "gross": bruto or 0, "net": neto or 0,
                        "stage": "cerrada" if estado in "DEF" else "en-curso", "month": cobro,
                    })

        if not listado:
            revisar.append({"kind": "fecha", "row": r, "property": prop,
                            "detail": "Sin fecha de ingreso: la columna de días publicada da 46.292."})
        if estado in "DEF" and not cobro:
            revisar.append({"kind": "fecha", "row": r, "property": prop,
                            "detail": f"Figura como {ESTADOS[estado][1].lower()} pero no tiene fechas ni Mes Impacto: no suma en ningún mes."})
        if nego is not None and abs(nego) > 50:
            revisar.append({"kind": "moneda", "row": r, "property": prop,
                            "detail": f"Publicada a {plata(publi)} y cerrada a {plata(cierre)}: parece mezclar pesos y dólares (da {nego:.1f} % de negociación). Va todo en dólares: el alquiler, al cambio del día de firma del contrato.".replace(f"{nego:.1f}", f"{nego:.1f}".replace(".", ","))})

    # Fórmulas pisadas a mano en la maestra: la que usa la mayoría de la columna es la regla, el resto se desvía.
    for k in ["Z", "AA"] + DINERO:
        patrones, textos = {}, {}
        for r in range(3, fs.max_row + 1):
            v = fs.cell(r, col(k)).value
            if v is None or not fs.cell(r, 2).value:
                continue
            texto = re.sub(rf"(?<=[A-Z]){r}\b", "#", str(v)).strip()
            # Se agrupa por la fórmula sin paréntesis de más, pero el aviso muestra la que está escrita.
            patrones.setdefault(sin_parentesis_de_mas(texto), []).append(r)
            textos[r] = texto
        if len(patrones) > 1:
            regla = max(patrones, key=lambda p: len(patrones[p]))
            for p, filas in patrones.items():
                if p == regla:
                    continue
                for r in filas:
                    revisar.append({"kind": "formula", "row": r, "property": fs.cell(r, 2).value, "column": k,
                                    "detail": f"La columna {ws.cell(2, col(k)).value} usa {formula_legible(textos[r], r)} en vez de la fórmula del resto."})

    # Las copias por agente contra la maestra: misma operación, montos distintos.
    for hoja in [n for n in wb.sheetnames if n not in NO_SON_COPIAS]:
        h = wb[hoja]
        # Una copia por agente tiene los mismos encabezados que la maestra; las pestañas nuevas
        # (Producción, Usuarios…) no, y compararlas daría diferencias que no existen.
        if str(h.cell(2, 2).value or "").strip() != "Propiedad":
            continue
        for r in range(3, h.max_row + 1):
            cierre = num(h.cell(r, col("W")).value)
            if not cierre or not h.cell(r, 2).value:
                continue
            candidatas = [o for o in operaciones if o["closePrice"] == cierre]
            if not candidatas:
                continue
            campos = {"AB": "gross", "AE": "producerNet", "AF": "agentNet", "AH": "leaderNet", "AI": "management", "AJ": "brea"}
            def diferencias(o):
                return [(k, num(h.cell(r, col(k)).value) or 0, o[f] or 0) for k, f in campos.items()
                        if abs((num(h.cell(r, col(k)).value) or 0) - (o[f] or 0)) > 1]
            o = min(candidatas, key=lambda o: len(diferencias(o)))
            for k, copia, maestra in diferencias(o):
                revisar.append({"kind": "copia", "row": o["row"], "property": o["property"], "sheet": hoja,
                                "detail": f"{ws.cell(2, col(k)).value}: {plata(copia)} en la hoja {hoja} y {plata(maestra)} en la maestra."})

    revisar = [i for i in revisar if i["row"] not in tres_arroyos and not (i["kind"] == "formula" and i["row"] in bolsas)]

    categorias = categorias_de_la_planilla(wb) or (leer_categorias(ruta_categorias) if ruta_categorias else [])
    # Cecilia es la gerente y Gonza el dueño: su % no sale de una categoría.
    # Gerente, dueño y externos no tienen categoría: su % no sale de la tabla.
    sin_categoria = {p for p, rol in directorio.roles.items() if rol.lower() in ("gerente", "dueño", "externo")}
    personas = [{"id": p, "name": NOMBRES[p], "rate": tasas.get(p), "role": directorio.roles.get(p),
                 "category": None if p in sin_categoria else categoria_de(tasas.get(p), categorias)} for p in NOMBRES]

    if categorias:
        # Confirmado con Gonzalo (23/10): la oficina se queda siempre el 35 % y el líder cobra la
        # diferencia con lo del productor (65 % menos su %). Se marca toda fila que no lo cumpla.
        for o in operaciones:
            tasa = tasas.get(o["producer"])
            if not o["leader"] or not o["producerGross"] or not isinstance(tasa, (int, float)):
                continue
            esperado = o["producerGross"] * max(0, 0.65 - tasa)
            if abs((o["leaderNet"] or 0) - esperado) > 1:
                revisar.append({"kind": "regla", "row": o["row"], "property": o["property"],
                    "detail": f"Líder {NOMBRES.get(o['leader'], o['leader'])}: cobra US$ {plata(o['leaderNet'] or 0)}. "
                              f"Con la regla (65 % menos el {round(tasa * 100)} % del productor) serían US$ {plata(esperado)}."})
        for pers in personas:
            if pers["id"] not in sin_categoria and pers["rate"] and not pers["category"]:
                revisar.append({"kind": "regla", "row": 0, "property": f"COMISIONES · {pers['name']}",
                    "detail": f"{pers['name']} figura con {round(pers['rate'] * 100)} %, que no está en la tabla de categorías (la más alta, Experto, es {round(max(c['agentRate'] for c in categorias) * 100)} %)."})
    # Nombres de la hoja 2026 que no están en Personas: el tablero los muestra, pero Producción no
    # (lista solo Personas) y no tienen %. Casi siempre es un typo o alguien nuevo sin cargar.
    for pid, nombre in directorio.desconocidos.items():
        filas = [o["row"] for o in operaciones if pid in (o["producer"], o["agent"], o["leader"])]
        AVISOS.append({"kind": "planilla", "row": filas[0] if filas else 0, "property": f"Personas · {nombre}",
                       "detail": f"«{nombre}» aparece en la hoja 2026 (filas {', '.join(map(str, filas))}) y no está en Personas: no figura en Producción ni tiene %."})

    costos, cargados = {}, {}
    for f in tabla(wb, "Costos", necesarias=("Mes", "Rubro", "Monto (ARS)", "Monto (US$)")):
        cuando, usd, ars = mes(f.get("Mes")), f.get("Monto (US$)"), f.get("Monto (ARS)")
        rubro = str(f.get("Rubro") or "").strip()
        donde = f"Costos · {rubro or 'sin rubro'}"
        if not cuando:
            if usd not in (None, "") or ars not in (None, ""):
                AVISOS.append({"kind": "planilla", "row": 0, "property": donde,
                               "detail": f"Fila {f['_fila']}: tiene monto pero la columna Mes no es una fecha. No suma en ningún mes."})
            continue
        if isinstance(usd, (int, float)) and not isinstance(usd, bool):
            costos[cuando] = costos.get(cuando, 0) + usd
            repetido = (cuando, sin_tildes(rubro).lower(), round(usd, 2))
            if repetido in cargados:
                AVISOS.append({"kind": "planilla", "row": 0, "property": donde,
                               "detail": f"Fila {f['_fila']}: mismo rubro, mes y monto que la fila {cargados[repetido]}. ¿Está cargado dos veces?"})
            cargados.setdefault(repetido, f["_fila"])
        elif ars not in (None, ""):
            AVISOS.append({"kind": "planilla", "row": 0, "property": donde,
                           "detail": f"Fila {f['_fila']}: tiene el monto en pesos pero no en dólares ({usd or 'vacío'}). Falta el dólar del mes en la columna D."})

    precios = [{"property": str(f["Propiedad"]).strip(), "date": dia(f.get("Fecha")), "price": num(f.get("Precio publicado (US$)"))}
               for f in tabla(wb, "Precios", necesarias=("Propiedad", "Fecha", "Precio publicado (US$)")) if f.get("Propiedad") and isinstance(f.get("Precio publicado (US$)"), (int, float))]
    propiedades = {o["property"].lower() for o in operaciones}
    for p in precios:
        if p["property"].lower() not in propiedades:
            AVISOS.append({"kind": "planilla", "row": 0, "property": f"Precios · {p['property']}",
                           "detail": "No coincide con ninguna propiedad de la hoja 2026: el cambio de precio no se ve en ninguna operación."})

    # Lo que el equipo explicó en la columna NOTA de Para revisar viaja con el aviso. Los números de fila
    # cambian cuando se agregan filas, así que se compara el texto sin números.
    clave = lambda prop, detalle: (str(prop).strip().lower(), re.sub(r"\d+", "#", str(detalle)))
    filas_nota = [f for f in tabla(wb, "Para revisar", 2, necesarias=("Propiedad", "Qué pasa", "NOTA")) if f.get("NOTA")]
    notas = {clave(f.get("Propiedad"), f.get("Qué pasa")): str(f["NOTA"]).strip() for f in filas_nota}
    usadas = set()
    for i in revisar:
        if clave(i["property"], i["detail"]) in notas:
            i["note"] = notas[clave(i["property"], i["detail"])]
            usadas.add(clave(i["property"], i["detail"]))
    sueltas = [f["_fila"] for f in filas_nota if clave(f.get("Propiedad"), f.get("Qué pasa")) not in usadas]
    if sueltas:
        AVISOS.append({"kind": "planilla", "row": 0, "property": "Para revisar",
                       "detail": f"{len(sueltas)} notas ya no corresponden a ningún aviso (filas {', '.join(map(str, sueltas))}): "
                                 "o se resolvió y se pueden borrar, o cambió el texto de «Qué pasa» y hay que copiar el nuevo."})

    # Los accesos van a un archivo aparte, nunca dentro de tablero.json: ese va entero en el paquete del admin.
    ACCESOS = {"ve todo": "admin", "líder": "lider", "lider": "lider", "agente": "agente"}
    accesos, usuarios_vistos = [], set()
    for f in tabla(wb, "Usuarios", necesarias=("Usuario", "Acceso", "Persona", "Equipo (si es líder)", "Contraseña nueva (opcional)")):
        usuario = str(f.get("Usuario") or "").strip().lower()
        acceso = str(f.get("Acceso") or "").strip()
        rol = ACCESOS.get(acceso.lower())
        if not usuario:
            continue
        donde = f"Usuarios · {usuario}"
        if not rol:
            AVISOS.append({"kind": "planilla", "row": 0, "property": donde,
                           "detail": f"El acceso «{acceso or 'vacío'}» no es Ve todo, Líder ni Agente: no entra al tablero."})
            continue
        if usuario in usuarios_vistos:
            AVISOS.append({"kind": "planilla", "row": 0, "property": donde, "detail": "Está dos veces: vale la primera fila."})
            continue
        usuarios_vistos.add(usuario)
        nombres_equipo = [n.strip() for n in str(f.get("Equipo (si es líder)") or "").split(",") if n.strip()]
        for n in [f.get("Persona"), *nombres_equipo]:
            if n and not directorio.conocida(n):
                AVISOS.append({"kind": "planilla", "row": 0, "property": donde,
                               "detail": f"«{n}» no está en Personas: con ese nombre no se le encuentran operaciones."})
        equipo = [directorio.persona(n) for n in nombres_equipo]
        accesos.append({"usuario": usuario, "rol": rol, "persona": directorio.persona(f.get("Persona")), "equipo": equipo,
                        "nombre": str(f.get("Nombre que se muestra") or "").strip() or None,
                        "contrasenaNueva": str(f.get("Contraseña nueva (opcional)") or "").strip() or None})
    # Siempre se escribe, aunque quede vacío: así cifrar.mjs corta con un mensaje claro en vez de usar
    # un accesos.json viejo. Va al lado de la salida, se llame como se llame.
    with open(os.path.join(os.path.dirname(os.path.abspath(salida)), "accesos.json"), "w") as a:
        json.dump(accesos, a, ensure_ascii=False, indent=1)

    AVISOS.extend(control_contra_la_planilla(wb, operaciones, participaciones, costos, NOMBRES))
    revisar.extend(AVISOS)

    json.dump({
        "generatedAt": HOY.isoformat(), "updatedAt": AHORA.isoformat(timespec="minutes"),
        "source": "2026 Operaciones Brea Propiedades · hoja 2026",
        "people": personas, "categories": categorias, "deals": operaciones, "shares": participaciones, "issues": revisar,
        "costs": [{"month": m, "usd": round(v, 2)} for m, v in sorted(costos.items())], "priceChanges": precios,
    }, open(salida, "w"), ensure_ascii=False, indent=1)
    print(f"{len(operaciones)} filas, {len(grupos)} operaciones, {len(participaciones)} participaciones, {len(revisar)} para revisar")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
